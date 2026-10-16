"""Exports the Gym CRM suite to an Excel workbook for testers.

Reads results/suite.json (written by build-report.ts) and results/results.json (written by
run-showcase.ts). Writes results/gym-crm-test-suite.xlsx. Usage:
    python showcase/fitness-crm/export_xlsx.py
"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
RESULTS = HERE / "results"
suite = json.loads((RESULTS / "suite.json").read_text(encoding="utf-8"))
results = json.loads((RESULTS / "results.json").read_text(encoding="utf-8"))

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F4E46")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # cells testers fill in
THIN = Side(style="thin", color="D0D7D4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
STATUS_FILL = {
    "PASSED": PatternFill("solid", fgColor="E3F3E8"),
    "FAILED": PatternFill("solid", fgColor="FBE7E5"),
    "INCONCLUSIVE": PatternFill("solid", fgColor="FDF0D9"),
    "BLOCKED": PatternFill("solid", fgColor="FDF0D9"),
}

# Latest result per case id, from each suite run (run 1 and run 2).
cases_out = {c["id"]: c for c in results["cases"]}
runs = []
for run in results["suiteRuns"]:
    by_tc = {}
    for item in run["items"]:
        if item.get("tc"):
            by_tc[item["tc"]] = item
    runs.append(by_tc)


def run_status(case_id, index):
    item = runs[index].get(case_id) if index < len(runs) else None
    if item is None:
        return ""
    return item["status"] if item.get("runId") else ""


def observed(case_id):
    item = runs[-1].get(case_id) if runs else None
    if not item or not item.get("detail"):
        return ""
    for a in item["detail"].get("assertions", []):
        if a.get("passed") is not True:
            return f'{a["stepId"]}: expected "{a.get("expected")}", observed "{str(a.get("actual"))[:200]}"'
    return ""


def case_type(case):
    if case["id"] in cases_out and not cases_out[case["id"]].get("caseId"):
        return "Not drafted"
    return "Automated" if case["automated"] else "Manual"


def style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = HEAD_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = BORDER


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()
summary = wb.active
summary.title = "Summary"

# --- One sheet per module
COLUMNS = ["ID", "Section", "Technique", "Priority", "Title", "Requirement ID", "Requirement", "Preconditions",
           "Steps (action → expected)", "Type", "Why manual", "Agent run 1", "Agent run 2", "Observed on failure",
           "Tester result", "Tester notes", "Tested on"]
WIDTHS = [11, 18, 20, 8, 40, 13, 50, 30, 60, 12, 40, 13, 13, 40, 13, 30, 12]
SHEETS = {"ACCESS": "Access", "LEAD": "Leads", "CON": "Customers & contracts", "SCHED": "Booking & sessions", "OPS": "Operations"}
module_sheets = []
result_dv_ranges = []

for module in suite["modules"]:
    name = SHEETS.get(module["id"], module["id"])[:31]
    ws = wb.create_sheet(name)
    module_sheets.append((module["name"], name))
    ws.append(COLUMNS)
    style_header(ws, 1, len(COLUMNS))
    for case in module["cases"]:
        steps = "\n".join(f'{i + 1}. {s["action"]} → {s["expected"]}' for i, s in enumerate(case["steps"]))
        ws.append([
            case["id"], case["section"], case["technique"], case["priority"], case["title"],
            case["requirement"]["id"], case["requirement"]["statement"], "\n".join(case.get("preconditions", [])),
            steps, case_type(case), case.get("manualReason") or "",
            run_status(case["id"], 0), run_status(case["id"], 1), observed(case["id"]), "", "", "",
        ])
        r = ws.max_row
        for col in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=r, column=col)
            cell.font = Font(name=FONT, size=10)
            cell.alignment = WRAP
            cell.border = BORDER
        for col in (12, 13):
            fill = STATUS_FILL.get(ws.cell(row=r, column=col).value)
            if fill:
                ws.cell(row=r, column=col).fill = fill
        for col in (15, 16, 17):
            ws.cell(row=r, column=col).fill = INPUT_FILL
    last = ws.max_row
    dv = DataValidation(type="list", formula1='"Pass,Fail,Blocked,Not run"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"O2:O{last}")
    ws.cell(row=1, column=15).comment = Comment("Fill in when running the case by hand: Pass, Fail, Blocked or Not run.", "QA agent")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"
    set_widths(ws, WIDTHS)

# --- Summary with live formulas over the module sheets
summary["A1"] = suite["name"]
summary["A1"].font = Font(name=FONT, bold=True, size=14)
ai = results["ai"]
summary["A2"] = ("AI: live " + ai["models"]["draft"]) if ai["live"] else "AI: keyless stand-in (drafts and readiness from a heuristic, not a model)"
summary["A3"] = f'Generated {results["generatedAt"][:10]} from the agent run against {results["target"]["url"]} ({results["target"]["product"]}).'
summary["A4"] = "Yellow cells on the module sheets are for testers: Tester result (Pass, Fail, Blocked, Not run), notes and date."
for r in (2, 3, 4):
    summary[f"A{r}"].font = Font(name=FONT, size=10, italic=True)

head = ["Module", "Sheet", "Cases", "Automated", "Manual", "Not drafted", "Agent: passed (run 2)", "Agent: failed (run 2)",
        "Agent: inconclusive (run 2)", "Tester: pass", "Tester: fail", "Tester: blocked"]
summary.append([])
summary.append(head)
hr = summary.max_row
style_header(summary, hr, len(head))
for module_name, sheet in module_sheets:
    q = f"'{sheet}'"
    summary.append([
        module_name, sheet,
        f"=COUNTA({q}!A:A)-1",
        f'=COUNTIF({q}!J:J,"Automated")',
        f'=COUNTIF({q}!J:J,"Manual")',
        f'=COUNTIF({q}!J:J,"Not drafted")',
        f'=COUNTIF({q}!M:M,"PASSED")',
        f'=COUNTIF({q}!M:M,"FAILED")',
        f'=COUNTIF({q}!M:M,"INCONCLUSIVE")',
        f'=COUNTIF({q}!O:O,"Pass")',
        f'=COUNTIF({q}!O:O,"Fail")',
        f'=COUNTIF({q}!O:O,"Blocked")',
    ])
first, last = hr + 1, summary.max_row
summary.append(["Total", ""] + [f"=SUM({get_column_letter(c)}{first}:{get_column_letter(c)}{last})" for c in range(3, len(head) + 1)])
for row in summary.iter_rows(min_row=first, max_row=summary.max_row, max_col=len(head)):
    for cell in row:
        cell.font = Font(name=FONT, size=10, bold=cell.row == summary.max_row)
        cell.border = BORDER
total_row = summary.max_row
summary.append(["Automation coverage", "", f"=IF(C{total_row}=0,0,D{total_row}/C{total_row})"])
summary.cell(row=summary.max_row, column=3).number_format = "0.0%"
summary.cell(row=summary.max_row, column=1).font = Font(name=FONT, size=10, bold=True)

summary.append([])
summary.append(["Technique", "Cases"])
th = summary.max_row
style_header(summary, th, 2)
techniques = sorted({c["technique"] for m in suite["modules"] for c in m["cases"]})
for t in techniques:
    parts = "+".join(f"COUNTIF('{s}'!C:C,\"{t}\")" for _, s in module_sheets)
    summary.append([t, f"={parts}"])
    for cell in summary[summary.max_row][:2]:
        cell.font = Font(name=FONT, size=10)
        cell.border = BORDER
set_widths(summary, [34, 22, 10, 11, 10, 12, 14, 14, 16, 12, 12, 13])

# --- Traceability
tr = wb.create_sheet("Traceability")
tr.append(["Requirement ID", "Statement (first case)", "Cases", "Agent results (run 2)"])
style_header(tr, 1, 4)
by_req = {}
for m in suite["modules"]:
    for c in m["cases"]:
        by_req.setdefault(c["requirement"]["id"], []).append(c)
for req_id in sorted(by_req):
    cs = by_req[req_id]
    tr.append([req_id, cs[0]["requirement"]["statement"], ", ".join(c["id"] for c in cs),
               ", ".join(f'{c["id"]}: {run_status(c["id"], 1) or case_type(c).lower()}' for c in cs)])
    for cell in tr[tr.max_row]:
        cell.font = Font(name=FONT, size=10)
        cell.alignment = WRAP
        cell.border = BORDER
tr.freeze_panes = "A2"
tr.auto_filter.ref = f"A1:D{tr.max_row}"
set_widths(tr, [16, 70, 40, 60])

# --- Test design artifacts
td = wb.create_sheet("Test design")
td["A1"] = "Test design artifacts: each table is the design a technique produced; the last column names the cases derived from it."
td["A1"].font = Font(name=FONT, italic=True, size=10)
for m in suite["modules"]:
    for d in m["designs"]:
        td.append([])
        td.append([f'{m["name"]}: {d["title"]}', d["technique"]])
        td.cell(row=td.max_row, column=1).font = Font(name=FONT, bold=True, size=11)
        td.cell(row=td.max_row, column=2).font = Font(name=FONT, italic=True, size=10)
        td.append(d["columns"])
        style_header(td, td.max_row, len(d["columns"]))
        for row in d["rows"]:
            td.append(row)
            for cell in td[td.max_row][: len(d["columns"])]:
                cell.font = Font(name=FONT, size=10)
                cell.alignment = WRAP
                cell.border = BORDER
set_widths(td, [34, 26, 24, 24, 22, 18, 18, 30, 12])

out = RESULTS / "gym-crm-test-suite.xlsx"
wb.save(out)
print(f"wrote {out}")
